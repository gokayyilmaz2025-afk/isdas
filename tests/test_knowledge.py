"""Real PDF/DOCX parsing, tenant retrieval and model context; synthetic HTTP only."""
import io,json,sqlite3,zipfile
from concurrent.futures import ThreadPoolExecutor
import httpx,pytest
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
from docx import Document
from test_api import clients,create_space
from backend import db,knowledge,providers,worker,ops


def pdf(pages=('Delivery takes 4 days.','Returns accepted within 21 days.'),encrypted=False):
    out=io.BytesIO();writer=PdfWriter()
    for text in pages:
        page=writer.add_blank_page(width=600,height=800)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream=DecodedStreamObject();stream.set_data(('BT /F1 12 Tf 40 750 Td ('+text+') Tj ET').encode('ascii'))
        page[NameObject('/Contents')]=writer._add_object(stream)
    if encrypted:writer.encrypt('synthetic-test-only')
    writer.write(out);return out.getvalue()


def docx():
    out=io.BytesIO();doc=Document();doc.add_paragraph('İade süremiz 21 gündür. Kargo hazırlığı 4 iş günü sürer.')
    table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Ürün';table.cell(0,1).text='Fiyat';table.cell(1,0).text='Çay';table.cell(1,1).text='45 TL'
    doc.save(out);return out.getvalue()


def upload(a,wid,raw=b'Kargo hazirligi 4 gun surer.',name='Bilgi.txt'):
    return a.post(f'/api/workspaces/{wid}/documents',files={'file':(name,raw,'application/octet-stream')})


@pytest.fixture
def library(clients,monkeypatch):
    a,b=clients;wid=create_space(a);other=create_space(b)
    monkeypatch.setenv('XAI_API_KEY','synthetic-document-test');monkeypatch.setenv('XAI_BASE_URL','https://document-model.example.test/v1')
    return a,b,wid,other


def test_real_pdf_and_docx_parse_with_source_positions_and_private_download(library):
    a,b,wid,other=library
    for name,raw,label,needle in [('Kosullar.pdf',pdf(),'Sayfa 2','21 days'),('Bilgiler.docx',docx(),'Tablo 1 · satır 2','45 TL')]:
        r=upload(a,wid,raw,name);assert r.status_code==200,r.text;did=r.json()['id']
        detail=a.get(f'/api/workspaces/{wid}/documents/{did}').json()
        assert any(c['label']==label and needle in c['body'] for c in detail['chunks'])
        download=a.get(f'/api/workspaces/{wid}/documents/{did}/download');assert download.content==raw
        assert download.headers['cache-control']=='no-store' and download.headers['content-type']=='application/octet-stream'
        for route in [f'/api/workspaces/{wid}/documents/{did}',f'/api/workspaces/{other}/documents/{did}',f'/api/workspaces/{wid}/documents/{did}/download']:
            assert b.get(route).status_code==404
    assert a.get(f'/api/workspaces/{wid}/documents').json()['bytes']>0


def test_rejects_scanned_encrypted_malformed_unsupported_and_large_files(library,monkeypatch):
    a,_,wid,_=library
    for raw,name,code in [(pdf(pages=('',)),'Scan.pdf',422),(pdf(encrypted=True),'Locked.pdf',422),(b'bad pdf','Bad.pdf',422),(b'bad zip','Bad.docx',422),(b'<script>bad</script>','Bad.html',400),(b'\xff\xff','Bad.txt',422)]:
        r=upload(a,wid,raw,name);assert r.status_code==code,r.text
    monkeypatch.setattr(knowledge,'MAX_UPLOAD',10)
    assert upload(a,wid,b'x'*11,'Large.txt').status_code==413
    assert not db.query('SELECT * FROM documents')


def test_pdf_empty_page_warning_and_large_text_are_not_silently_ignored(library):
    a,_,wid,_=library
    r=upload(a,wid,pdf(pages=('Delivery policy is four days.','')),'Partial.pdf');assert r.status_code==200
    assert any('sayfalar: 2' in text for text in r.json()['warnings'])
    huge=upload(a,wid,b'A'*200001,'Huge.txt');assert huge.status_code==422
    assert len(db.query('SELECT * FROM documents'))==1


def test_duplicate_upload_and_concurrent_quota_are_atomic(library,monkeypatch):
    a,_,wid,_=library;raw=b'Unique document content for atomic upload.'
    with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(lambda _:upload(a,wid,raw),range(2)))
    assert all(r.status_code==200 for r in results)
    assert len({r.json()['id'] for r in results})==1
    assert len(db.query('SELECT * FROM document_blobs'))==1
    monkeypatch.setattr(knowledge,'MAX_STORAGE',len(raw)+5)
    assert upload(a,wid,b'Different document content','Another.txt').status_code==409
    assert len(db.query('SELECT * FROM document_blobs'))==1


def test_search_is_workspace_scoped_turkish_and_does_not_execute_query_syntax(library):
    a,b,wid,other=library
    upload(a,wid,'İADE süresi 21 gündür. IŞIK mağazası.'.encode(),'Bilgi.txt')
    upload(b,other,'İade süresi 99 gündür. Gizli müşteri.'.encode(),'Gizli.txt')
    result=a.get(f'/api/workspaces/{wid}/document-search',params={'q':'iade'}).json()['items']
    assert len(result)==1 and '21' in result[0]['body'] and 'Gizli' not in str(result)
    assert knowledge.retrieve(wid,'ışık')[0]['name']=='Bilgi.txt'
    assert b.get(f'/api/workspaces/{wid}/document-search?q=iade').status_code==404
    assert a.get(f'/api/workspaces/{wid}/document-search',params={'q':'" OR * NOT (iade)'}).status_code==200


def test_disable_delete_scope_member_permissions_and_search_index_cleanup(library):
    a,b,wid,_=library;did=upload(a,wid).json()['id'];base=f'/api/workspaces/{wid}/documents/{did}'
    db.query('INSERT INTO memberships VALUES(?,?,?)',(wid,b.get('/api/me').json()['id'],'member'))
    assert b.get(base).status_code==200
    assert b.patch(base,json={'enabled':False}).status_code==403 and b.delete(base).status_code==403
    assert a.patch(base,json={'enabled':False}).status_code==200
    assert knowledge.retrieve(wid,'kargo')==[]
    with pytest.raises(knowledge.SelectionUnavailable):knowledge.retrieve(wid,'kargo',[did])
    assert a.patch(base,json={'enabled':True}).status_code==200
    assert knowledge.retrieve(wid,'kargo')
    assert a.delete(base).status_code==200
    assert not db.query('SELECT * FROM document_chunks') and not db.query('SELECT * FROM document_blobs')
    assert knowledge.retrieve(wid,'kargo')==[]


def test_selected_preview_is_bounded_and_does_not_claim_full_document(library):
    a,_,wid,_=library
    raw=('Birinci bölüm. '+'Uzun çalışma notları ve ayrıntılar. '*4000).encode()
    did=upload(a,wid,raw,'Uzun.txt').json()['id']
    sources=knowledge.retrieve(wid,'Özet',[did])
    assert len(sources)<=8 and sum(len(s['body']) for s in sources)<=12800
    assert {s['retrieval'] for s in sources}=={'selected_sample'}
    assert sources[-1]['ordinal']>sources[0]['ordinal']


def test_multiple_selected_documents_are_represented_with_extraction_warnings(library):
    a,_,wid,_=library;first=upload(a,wid,b'Delivery time is 4 days.','Policy.txt').json()['id'];second=upload(a,wid,docx(),'Marka.docx').json()['id']
    sources=knowledge.retrieve(wid,'Delivery',[first,second])
    assert {s['document_id'] for s in sources}=={first,second}
    word=next(s for s in sources if s['document_id']==second)
    assert word['retrieval']=='selected_sample' and word['warnings']


def configure_model(monkeypatch,callback):
    original=httpx.Client
    def reply(request):return httpx.Response(200,json=callback(json.loads(request.content)))
    monkeypatch.setattr(providers.httpx,'Client',lambda *a,**k:original(transport=httpx.MockTransport(reply)))


def result(text):return {'id':'document-test','status':'completed','output':[{'type':'message','content':[{'type':'output_text','text':text}]}],'usage':{'cost_in_usd_ticks':1000000}}


def submit(a,wid,docs=None,key=None):
    return a.post(f'/api/workspaces/{wid}/jobs',json={'prompt':'İade süresi ne kadar?','agent_id':'guide','idempotency_key':key or db.uid(),'start_conversation':True,'document_ids':docs or []})


def complete():
    job=worker.claim();assert job;worker.process(job);return db.query('SELECT * FROM jobs WHERE id=?',(job['id'],),one=True)


def test_provider_gets_only_selected_own_sources_and_citations_are_inspectable(library,monkeypatch):
    a,b,wid,other=library;did=upload(a,wid,docx(),'Politika.docx').json()['id'];upload(b,other,b'Private unrelated secret','Private.txt')
    calls=[]
    def respond(body):calls.append(body);return result('TEST: İade süresi 21 gündür [K1].')
    configure_model(monkeypatch,respond)
    job=submit(a,wid,[did]).json();done=complete();assert done['status']=='completed'
    system=calls[0]['input'][0]['content'];assert 'document_excerpts' in system and '21 gündür' in system and 'Private unrelated' not in system
    sources=a.get(f"/api/workspaces/{wid}/jobs/{job['id']}/documents").json()['items'];assert sources[0]['cited'] is True and sources[0]['available'] is True
    detail=a.get(f"/api/workspaces/{wid}/conversations/{job['conversation_id']}").json();assert detail['turns'][0]['document_source_count']>=1
    assert b.get(f"/api/workspaces/{wid}/jobs/{job['id']}/documents").status_code==404
    a.delete(f'/api/workspaces/{wid}/documents/{did}')
    deleted=a.get(f"/api/workspaces/{wid}/jobs/{job['id']}/documents").json()['items'];assert deleted[0]['available'] is False and 'body' not in deleted[0]
    assert '21 gündür' in db.query('SELECT output FROM jobs WHERE id=?',(job['id'],),one=True)['output']


def test_explicit_selection_cannot_cross_workspace_or_change_on_replay(library):
    a,b,wid,other=library;did=upload(a,wid).json()['id'];foreign=upload(b,other).json()['id'];key=db.uid()
    assert submit(a,wid,[foreign]).status_code==404
    first=submit(a,wid,[did],key);assert first.status_code==200
    assert submit(a,wid,[did],key).json()['id']==first.json()['id']
    assert submit(a,wid,[],key).status_code==409


def test_removed_selected_document_before_call_releases_reservation(library,monkeypatch):
    a,_,wid,_=library;did=upload(a,wid).json()['id'];job=submit(a,wid,[did]).json()
    a.delete(f'/api/workspaces/{wid}/documents/{did}')
    configure_model(monkeypatch,lambda body:pytest.fail('Provider must not be called'))
    done=complete();assert done['status']=='failed' and 'Seçtiğin belge' in done['error']
    assert db.query('SELECT state FROM credit_ledger WHERE job_id=?',(job['id'],),one=True)['state']=='released'


def test_document_changed_during_call_discards_output_but_keeps_actual_cost(library,monkeypatch):
    a,_,wid,_=library;did=upload(a,wid,docx(),'Politika.docx').json()['id'];job=submit(a,wid,[did]).json()
    def respond(body):
        assert a.patch(f'/api/workspaces/{wid}/documents/{did}',json={'enabled':False}).status_code==200
        return result('TEST: Eski belge bilgisi [K1].')
    configure_model(monkeypatch,respond);done=complete()
    assert done['status']=='failed' and not done['output'] and 'belge kaldırıldı' in done['error']
    assert not db.query('SELECT * FROM job_document_sources') and not db.query('SELECT * FROM artifacts')
    assert db.query('SELECT actual_ticks FROM usage_ledger WHERE job_id=?',(job['id'],),one=True)['actual_ticks']==1000000


def test_unknown_citation_is_marked_unverified_and_receipt_retained(library,monkeypatch):
    a,_,wid,_=library;configure_model(monkeypatch,lambda body:result('TEST: İddia [K999].'))
    submit(a,wid);done=complete();assert done['status']=='completed'
    assert '[K999]' not in done['output'] and 'belge atfı doğrulanamadı' in done['output']
    assert json.loads(done['context_info'])['document_citation_warning'] is True


def test_parser_timeout_leaves_no_document_and_environment_has_no_keys(library,monkeypatch):
    a,_,wid,_=library;original=knowledge.subprocess.run;seen=[]
    def spy(*args,**kwargs):seen.append(kwargs['env']);return original(*args,**kwargs)
    monkeypatch.setattr(knowledge.subprocess,'run',spy);monkeypatch.setattr(knowledge,'PARSER_TIMEOUT',.001)
    assert upload(a,wid,pdf(),'Slow.pdf').status_code==422
    assert 'XAI_API_KEY' not in seen[0] and 'ENCRYPTION_KEY' not in seen[0]
    assert not db.query('SELECT * FROM documents')


def test_backup_roundtrip_preserves_documents_and_search_index(library,tmp_path,monkeypatch):
    from cryptography.fernet import Fernet
    a,_,wid,_=library;did=upload(a,wid,docx(),'Politika.docx').json()['id']
    monkeypatch.setenv('BACKUP_ENCRYPTION_KEY',Fernet.generate_key().decode())
    target=tmp_path/'documents.enc';ops.backup(db.DB_PATH,target);restored=tmp_path/'restored.sqlite3';ops.restore_candidate(target,restored)
    with sqlite3.connect(restored) as conn:
        assert conn.execute('SELECT count(*) FROM document_blobs').fetchone()[0]==1
        assert conn.execute("SELECT count(*) FROM document_search WHERE document_search MATCH 'iade'").fetchone()[0]>=1
        assert conn.execute('SELECT id FROM documents').fetchone()[0]==did
